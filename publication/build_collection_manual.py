"""Build the operational manual from its versioned Markdown source."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]

def clean(text):
    text=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text)
    return text.replace('`','').replace('**','')

def main():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    for element in list(doc.styles.element.iter(qn('w:pBdr'))):element.getparent().remove(element)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(23)
    doc.styles['Heading 1'].font.size=Pt(15)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    lines=(ROOT/'RECOLECCION_HISTORICA.md').read_text().splitlines();index=0;code=False
    while index<len(lines):
        line=lines[index];index+=1
        if line.startswith('```'):code=not code;continue
        if not line.strip():continue
        if code:
            p=doc.add_paragraph(line);p.paragraph_format.space_after=Pt(2)
            for run in p.runs:run.font.name='Consolas';run.font.size=Pt(9)
            continue
        if line.startswith('# '):doc.add_paragraph(clean(line[2:]),'Title');continue
        if line.startswith('## '):doc.add_heading(clean(line[3:]),1);continue
        if line.startswith('|'):
            rows=[line]
            while index<len(lines) and lines[index].startswith('|'):rows.append(lines[index]);index+=1
            values=[[clean(x.strip()) for x in row.strip('|').split('|')] for row in rows if not re.match(r'^\|[\s:|\-]+\|$',row)]
            table=doc.add_table(rows=0,cols=len(values[0]));table.autofit=False
            widths=[1.8,4.7] if len(values[0])==2 else [1.45,2.5,2.55]
            for col,width in zip(table.columns,widths):col.width=Inches(width)
            for r,value in enumerate(values):
                cells=table.add_row().cells
                for cell,text,width in zip(cells,value,widths):
                    cell.width=Inches(width);cell.text=text;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                    for name in ['top','left','bottom','right']:
                        edge=OxmlElement('w:'+name);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                    pr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for name in ['top','bottom','left','right']:
                        edge=OxmlElement('w:'+name);edge.set(qn('w:w'),'90');edge.set(qn('w:type'),'dxa');margins.append(edge)
                    pr.append(margins)
                    if r==0:
                        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E6EDF3');pr.append(shd)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3)
                        for run in p.runs:run.font.size=Pt(10);run.bold=r==0
                if r==0:
                    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                no_split=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(no_split)
            doc.add_paragraph();continue
        numbered=bool(re.match(r'^\d+\.\s+',line))
        line=re.sub(r'^\d+\.\s+', '',line)
        doc.add_paragraph(clean(line),'List Number' if numbered else None)
    doc.core_properties.title='Recolección histórica recuperable de SIAN'
    doc.core_properties.subject='Coordinador de ejecución versión 3'
    doc.core_properties.author='SIAN'
    out=ROOT/'publication/SIAN_Recoleccion_Historica_2026-10-07.docx';doc.save(out)
    print(out)
if __name__=='__main__':main()
